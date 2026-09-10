from io import BytesIO
import json
from threading import Event
from uuid import uuid4

from docx import Document as Word
from openpyxl import Workbook
from sqlalchemy import select

from tests.conftest import auth_headers, PASSWORD
from app.core.config import get_settings
from app.db.session import SessionLocal
from app.models.message import Message
from app.models.document import Document
from app.services.chunking_service import ChunkingService
from app.services.extraction_service import ExtractionService, ExtractionError
from app.services.document_service import DocumentService
from app.services.query_analysis_service import QuestionAnalysisService
from app.services.rag_service import RAGService, GenerationCancelled
from app.services.vector_store_service import RetrievedChunk
from app.services import generation_service


def upload(client, headers, name="politica.txt", text="A política exige aprovação da gerência. O prazo de reembolso é de 30 dias."):
    result = client.post("/api/documents", headers=headers, files={"file":(name,text.encode(),"text/plain")})
    assert result.status_code == 201, result.text
    return result.json()


def chat(client, headers, ids=None):
    result = client.post("/api/chats", headers=headers, json={"document_ids":ids})
    assert result.status_code == 201, result.text
    return result.json()["id"]


def test_token_chunking_covers_end_and_location():
    text = " ".join("palavra"+str(i) for i in range(300)) + " ZX-900"
    counter = lambda value: len(value.split()) + 2
    chunks = ChunkingService(1200,150,token_counter=counter,token_limit=40,token_overlap=5).split([{"text":text,"page":7,"location":"Seção final"}])
    assert all(counter(c.text) <= 40 for c in chunks)
    assert all(c.location == "Seção final" and c.page == 7 for c in chunks)
    assert "ZX-900" in chunks[-1].text
    assert all("palavra"+str(i) in " ".join(c.text for c in chunks) for i in range(300))


def test_xlsx_preserves_empty_columns(tmp_path):
    workbook = Workbook()
    workbook.active.append(["Nome","Prazo","Valor"])
    workbook.active.append(["Ana",None,1234])
    path=tmp_path/"colunas.xlsx";workbook.save(path)
    sections=ExtractionService(get_settings()).extract(path,"XLSX")
    row=next(s for s in sections if "Ana" in s["text"])
    assert "Prazo: (vazio" in row["text"]
    assert "Valor: 1234" in row["text"]
    assert "2" in row["location"]


def test_docx_keeps_tables_in_reading_order(tmp_path):
    doc=Word();doc.add_paragraph("Antes da tabela")
    table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text="Código";table.cell(0,1).text="Valor"
    table.cell(1,0).text="ZX-900";table.cell(1,1).text="500"
    doc.add_paragraph("Depois da tabela")
    path=tmp_path/"tabela.docx";doc.save(path)
    content="\n".join(s["text"] for s in ExtractionService(get_settings()).extract(path,"DOCX"))
    assert content.index("Antes") < content.index("ZX-900") < content.index("Depois")


def test_followup_and_explicit_topic_are_distinguished():
    assert QuestionAnalysisService.analyze("Qual é o prazo?",True).uses_history
    assert not QuestionAnalysisService.analyze("Pode explicar o contrato ZX-900?",True).uses_history
    result=QuestionAnalysisService.analyze("Resuma em dois tópicos",True)
    assert result.item_limit == 2 and "2" in result.task_instruction


def test_citations_preserve_same_page_and_multiple_pages():
    chunks=[RetrievedChunk(str(i),text,1,"manual.pdf",page,1,citation=i)
            for i,text,page in [(1,"Aprovação em 30 dias",1),(2,"Valor máximo 500",1),(3,"Exceção para gestores",2)]]
    assert len(RAGService._supporting_sources(chunks,"Prazo [1]. Valor [2]. Exceção [3].")) == 3


def test_duplicate_version_archive_preview_and_scope(client):
    h=auth_headers(client);other=auth_headers(client,"outro@aiup.local")
    first=upload(client,h)
    duplicate=client.post("/api/documents",headers=h,files={"file":("copia.txt",b"A pol\xc3\xadtica exige aprova\xc3\xa7\xc3\xa3o da ger\xc3\xaancia. O prazo de reembolso \xc3\xa9 de 30 dias.","text/plain")})
    assert duplicate.status_code == 409
    second=upload(client,h,text="Nova política: prazo de reembolso de 45 dias.")
    assert second["version"] == 2
    assert client.get("/api/documents/"+str(first["id"]),headers=h).json()["archived"]
    assert client.get(f"/api/documents/{second['id']}/preview",headers=h).json()["sections"]
    assert client.get(f"/api/documents/{second['id']}/file",headers=other).status_code == 404
    assert client.post("/api/chats",headers=other,json={"document_ids":[second["id"]]}).status_code == 404
    cid=chat(client,h,[first["id"]])
    result=client.post(f"/api/chats/{cid}/messages",headers=h,json={"message":"Qual é o prazo de reembolso?"})
    assert result.status_code == 200
    assert result.json()["sources"] and all(s["document_id"] == first["id"] for s in result.json()["sources"])


def test_failed_reprocess_keeps_previous_index(client,monkeypatch):
    h=auth_headers(client);item=upload(client,h)
    with SessionLocal() as db:
        doc=db.get(Document,item["id"]);previous=doc.index_generation
        def fail(*args): raise ExtractionError("PDF indisponível")
        service=DocumentService(get_settings());monkeypatch.setattr(service.extractor,"extract",fail)
        service.process(db,doc);db.refresh(doc)
        assert doc.status.value == "READY" and doc.index_generation == previous
        assert "indisponível" in doc.error_message


def test_last_admin_null_name_and_password_revocation(client):
    h=auth_headers(client,"admin@aiup.local")
    me=client.get("/api/auth/me",headers=h).json()
    assert client.patch(f"/api/users/{me['id']}",headers=h,json={"role":"USER"}).status_code == 409
    assert client.patch(f"/api/users/{me['id']}",headers=h,json={"name":None}).status_code == 422
    result=client.post("/api/auth/password",headers=h,json={"current_password":PASSWORD,"new_password":"New@Password123"})
    assert result.status_code == 200
    assert client.get("/api/auth/me",headers=h).status_code == 401
    assert client.post("/api/auth/login",json={"email":me["email"],"password":"New@Password123"}).status_code == 200


def test_login_rate_limit(client):
    for _ in range(get_settings().login_attempt_limit):
        assert client.post("/api/auth/login",json={"email":"usuario@aiup.local","password":"wrong"}).status_code == 401
    limited = client.post("/api/auth/login",json={"email":"usuario@aiup.local","password":"wrong"})
    assert limited.status_code == 429
    assert limited.json()["error"] == "RATE_LIMITED"
    assert limited.headers["Retry-After"] == str(get_settings().login_attempt_window_seconds)


def test_job_idempotency_retry_export_feedback_and_sources_survive_delete(client,monkeypatch):
    # Run the worker synchronously to verify its persistence and SSE replay deterministically.
    def submit(key):
        state=generation_service.RunningGeneration();generation_service._states[key]=state
        generation_service.run(key,state)
    monkeypatch.setattr(generation_service,"submit",submit)
    h=auth_headers(client);doc=upload(client,h);cid=chat(client,h,[doc["id"]])
    payload={"message":"Qual é o prazo de reembolso?","request_id":str(uuid4())}
    first=client.post(f"/api/chats/{cid}/generations",headers=h,json=payload)
    assert first.status_code == 202,first.text
    replay=client.post(f"/api/chats/{cid}/generations",headers=h,json=payload).json()
    assert replay["status"] == "done",replay
    mid=replay["result"]["message_id"]
    assert client.get(f"/api/chats/{cid}",headers=h).json()["total_messages"] == 2
    events=client.get(f"/api/chats/{cid}/generations/{replay['id']}/events",headers=h)
    assert '"type": "done"' in events.text
    retry={**payload,"request_id":str(uuid4()),"retry_message_id":mid}
    assert client.post(f"/api/chats/{cid}/generations",headers=h,json=retry).status_code == 202
    assert client.get(f"/api/chats/{cid}",headers=h).json()["total_messages"] == 2
    assert client.patch(f"/api/chats/{cid}/messages/{mid}/feedback",headers=h,json={"feedback":-1,"reason":"Fonte inadequada"}).status_code == 200
    exported=client.get(f"/api/chats/{cid}/export?format=json",headers=h)
    assert exported.status_code == 200 and "Fonte inadequada" in exported.text
    assert client.delete(f"/api/documents/{doc['id']}",headers=h).status_code == 204
    sources=client.get(f"/api/chats/{cid}",headers=h).json()["messages"][-1]["sources"]
    assert sources and sources[0]["document"] == "politica.txt" and sources[0]["document_id"] is None


def test_cancelled_generation_does_not_save_messages(client,monkeypatch):
    h=auth_headers(client);cid=chat(client,h)
    monkeypatch.setattr(generation_service,"submit",lambda key:None)
    job=client.post(f"/api/chats/{cid}/generations",headers=h,json={"message":"Qual prazo?","request_id":str(uuid4())}).json()
    assert client.post(f"/api/chats/{cid}/generations/{job['id']}/cancel",headers=h).json()["status"] == "cancelled"
    generation_service.run(job["id"],generation_service.RunningGeneration())
    assert client.get(f"/api/chats/{cid}",headers=h).json()["total_messages"] == 0


def test_document_filters_and_pdf_render_authorization(client):
    import pymupdf
    h=auth_headers(client);other=auth_headers(client,"outro@aiup.local")
    pdf=pymupdf.open();page=pdf.new_page();page.insert_text((72,72),"Regra de demonstracao: prazo de 30 dias.")
    item=client.post("/api/documents",headers=h,files={"file":("paginas.pdf",pdf.tobytes(),"application/pdf")}).json();pdf.close()
    assert client.patch(f"/api/documents/{item['id']}",headers=h,json={"category":"Contratos","archived":True}).status_code == 200
    assert client.get("/api/documents?category=Contratos&archived=true",headers=h).json()[0]["id"] == item["id"]
    assert client.get("/api/documents?archived=false",headers=h).json() == []
    result=client.get(f"/api/documents/{item['id']}/pages/1",headers=h)
    assert result.status_code == 200 and result.content.startswith(b"\x89PNG")
    assert client.get(f"/api/documents/{item['id']}/pages/99",headers=h).status_code == 404
    assert client.get(f"/api/documents/{item['id']}/pages/1",headers=other).status_code == 404


def test_cancellation_during_generation_rolls_back(client,monkeypatch):
    h=auth_headers(client);doc=upload(client,h);cid=chat(client,h,[doc["id"]])
    stopped=Event()
    class Provider:
        def generate_stream(self,*args):
            yield "Parcial"
            stopped.set()
            yield " nunca deve persistir"
    monkeypatch.setattr("app.services.rag_service.get_llm_provider",lambda _:Provider())
    from app.models.chat import Chat
    from app.models.user import User
    import pytest
    with SessionLocal() as db:
        c=db.get(Chat,cid)
        with pytest.raises(GenerationCancelled):
            RAGService(get_settings()).ask(db,c,db.get(User,c.user_id),"Qual é o prazo?",emit=lambda *args:None,cancelled=stopped.is_set)
        db.rollback()
        assert list(db.scalars(select(Message).where(Message.chat_id==cid))) == []


def test_reduction_accounts_for_many_short_section_labels():
    chunks=[RetrievedChunk(str(i),"Dado "+str(i),1,"manual-muito-longo.pdf",1,1,citation=i) for i in range(300)]
    calls=[]
    class Provider:
        def generate(self,question,context,*args):
            calls.append(context)
            return "Síntese [1]."
    result=RAGService(get_settings())._reduce(Provider(),chunks,"Resuma",lambda stage:None)
    assert calls and sum(len(group) for group in calls) == 300
    assert RAGService._context_size(result) <= 18000


def test_retrieval_hit_requires_expected_page():
    from scripts.evaluate_ai import retrieval_hit
    wrong=RetrievedChunk("wrong","prazo",1,"regulamento.pdf",1,1)
    assert not retrieval_hit([wrong],{"expected_document_id":1,"expected_page":2})
