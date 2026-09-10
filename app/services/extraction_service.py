from io import BytesIO
from pathlib import Path
import shutil

import pymupdf
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from PIL import Image
import pytesseract

from app.core.config import Settings


class ExtractionError(RuntimeError):
    pass


class ExtractionService:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self.warnings: list[str] = []
        if settings.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd

    def extract(self, path: Path, file_type: str) -> list[dict]:
        self.warnings = []
        handlers = {
            "PDF": self._extract_pdf,
            "TXT": self._extract_txt,
            "DOCX": self._extract_docx,
            "XLSX": self._extract_xlsx,
            "IMAGE": self._extract_image,
        }
        try:
            sections = handlers[file_type.upper()](path)
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError(f"Falha ao extrair {file_type}: {exc}") from exc
        if not any(section.get("text", "").strip() for section in sections):
            raise ExtractionError("Nenhum texto foi encontrado no documento.")
        if sum(len(section.get("text", "")) for section in sections) > self.settings.max_extracted_chars:
            raise ExtractionError("O documento contém texto demais. Divida-o em arquivos menores.")
        return sections

    @staticmethod
    def _extract_txt(path: Path) -> list[dict]:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")
        return [{"text": text, "page": None}]

    def _extract_pdf(self, path: Path) -> list[dict]:
        sections: list[dict] = []
        with pymupdf.open(path) as pdf:
            if pdf.needs_pass:
                raise ExtractionError("O PDF está protegido por senha. Envie uma cópia desbloqueada.")
            if len(pdf) > self.settings.max_document_pages:
                raise ExtractionError("O PDF excede o limite de páginas. Divida o arquivo.")
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()
                if len(text) < 40:
                    try:
                        text = self._ocr_pdf_page(page) or text
                    except Exception:
                        self.warnings.append(f"Página {page_number}: OCR indisponível; confira o conteúdo original.")
                sections.append({"text": text, "page": page_number})
        return sections

    def _ocr_pdf_page(self, page: pymupdf.Page) -> str:
        self._ensure_tesseract()
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
        image = Image.open(BytesIO(pixmap.tobytes("png")))
        return pytesseract.image_to_string(image, lang="por", timeout=45).strip()

    @staticmethod
    def _extract_docx(path: Path) -> list[dict]:
        document = DocxDocument(path)
        sections = []
        heading = "Documento"
        paragraph_number = table_number = 0
        for child in document.element.body:
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                paragraph_number += 1
                if paragraph.text.strip():
                    if paragraph.style and paragraph.style.name.startswith("Heading"):
                        heading = paragraph.text.strip()
                    sections.append({"text": paragraph.text, "page": None,
                                     "location": f"{heading[:150]} · parágrafo {paragraph_number}"})
            elif child.tag.endswith("}tbl"):
                table_number += 1
                table = Table(child, document)
                headers = [cell.text.strip() or f"Coluna {i+1}" for i, cell in enumerate(table.rows[0].cells)] if table.rows else []
                for number, row in enumerate(table.rows, 1):
                    values = [cell.text.strip() or "(vazio)" for cell in row.cells]
                    sections.append({"text": " | ".join(f"{headers[i] if i < len(headers) else i+1}: {value}" for i, value in enumerate(values)),
                                     "page": None, "location": f"Tabela {table_number} · linha {number}"})
        return sections

    @staticmethod
    def _extract_xlsx(path: Path) -> list[dict]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sections: list[dict] = []
        try:
            for sheet in workbook.worksheets:
                if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > 500000:
                    raise ExtractionError("Planilha muito grande. Exporte somente a região com dados.")
                headers = None
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if not any(value is not None for value in row):
                        continue
                    if headers is None:
                        headers = [str(value).strip() if value is not None else f"Coluna {i+1}" for i, value in enumerate(row)]
                    values = [str(value).strip() if value is not None else "(vazio/sem valor calculado)" for value in row]
                    sections.append({"text": f"Planilha {sheet.title}, linha {number}: " + " | ".join(f"{headers[i] if i < len(headers) else f'Coluna {i+1}'}: {value}" for i, value in enumerate(values)),
                                     "page": None, "sheet": sheet.title, "location": f"Linha {number}"})
        finally:
            workbook.close()
        return sections

    def _extract_image(self, path: Path) -> list[dict]:
        self._ensure_tesseract()
        with Image.open(path) as image:
            text = pytesseract.image_to_string(image, lang="por", timeout=45)
        return [{"text": text, "page": None}]

    def _ensure_tesseract(self) -> None:
        configured = self.settings.tesseract_cmd
        if configured and Path(configured).is_file():
            return
        if shutil.which("tesseract"):
            return
        raise ExtractionError(
            "Tesseract não encontrado. Instale-o ou configure TESSERACT_CMD no arquivo .env."
        )
